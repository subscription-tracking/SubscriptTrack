from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

OUT = Path(r"D:\SubscriptTrack\outputs\reports\SubscriptTrack_Teknik_ve_Teknoloji_Raporu.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY, BLUE, LIGHT, GRAY = "17365D", "2F75B5", "EAF2F8", "5B6573"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(t):
    tblPr = t._tbl.tblPr
    bd = tblPr.first_child_found_in("w:tblBorders")
    if bd is None:
        bd = OxmlElement("w:tblBorders"); tblPr.append(bd)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = bd.find(qn("w:"+edge))
        if e is None: e = OxmlElement("w:"+edge); bd.append(e)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "D9D9D9")

def put(cell, text, bold=False, color="000000", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def make_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)
    for i, h in enumerate(headers):
        put(t.rows[0].cells[i], h, True, "FFFFFF", 9); shade(t.rows[0].cells[i], NAVY)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            put(cells[i], value)
            if n % 2: shade(cells[i], "F6F9FC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(text); return p

def code(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("2F3B4A")

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10); styles["Normal"].font.color.rgb = RGBColor(35,42,52)
styles["Normal"].paragraph_format.space_after = Pt(6); styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size, color in [("Title",28,NAVY),("Heading 1",18,NAVY),("Heading 2",13,BLUE),("Heading 3",11,NAVY)]:
    s = styles[name]; s.font.name = "Aptos Display"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
header = sec.header.paragraphs[0]; header.text = "SUBSCRIPTRACK  |  TEKNİK ÜRÜN RAPORU"; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GRAY); r.bold = True
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("SubscriptTrack • Teknik ve Teknolojik Yapı"); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GRAY)

p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("SubscriptTrack Teknik ve Teknolojik Yapı Raporu")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Mobil abonelik takip ürünü • Mimari, kod yapısı, veri akışları ve canlı durum"); r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(GRAY)
doc.add_paragraph("")
make_table(doc, ["Rapor bilgisi","Değer"], [
    ("Rapor tarihi","22 Eylül 2026"), ("Kapsam","Flutter mobil uygulama, Supabase backend, GitHub Pages ve operasyonel akışlar"),
    ("Kaynak durum","main branch • commit 6e116b2 • release APK 59,9 MB"), ("Canlı Supabase","tdbljrojcmyjwchawfif projesi")
], [2.0,4.8])
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Bu rapor, ürünün teknik yapısını ve mevcut uygulama durumunu tek bir referansta toplar.").italic = True
doc.add_page_break()

doc.add_heading("Yönetici Özeti", 1)
doc.add_paragraph("SubscriptTrack, banka veya kart hesabı bağlamadan dijital abonelikleri takip eden, iOS ve Android için Flutter ile geliştirilmiş mobil üründür. Mobil uygulama kimlik doğrulama, abonelik yaşam döngüsü, dashboard, takvim, tasarruf, ödeme etiketi, yerel bildirim, CSV dışa aktarma ve çevrimdışı son veri görüntüleme yeteneklerini bir araya getirir. Supabase; Auth, PostgreSQL, Row Level Security, Realtime ve Edge Functions katmanlarını sağlar.")
doc.add_paragraph("Mevcut sürümün kod ve backend omurgası çalışır durumdadır. Bildirimler uzak push değil, cihaz üzerinde planlanan yerel bildirimlerdir. Web yüzü tam dashboard değil, Flutter web önizlemesi ve auth dönüş yüzüdür. Gerçek e-posta teslimi, deep-link ve iki cihazlı Realtime kabulü gerçek hesap ve cihazla sonlandırılmalıdır.")
make_table(doc, ["Alan","Durum","Kanıt / sınır"], [
    ("Mobil istemci","Tamamlandı","Flutter analiz/testleri başarılı; release APK üretildi"),
    ("Supabase Auth","Kod ve canlı yapı hazır","Mail sağlayıcısı ve gerçek e-posta testi dış bağımlılık"),
    ("Abonelik CRUD","Canlı Supabase RPC + RLS","Authenticated smoke script hazır"),
    ("Bildirim","Yerel cihaz planlaması","FCM/APNs kullanılmıyor; cihaz izni gerekir"),
    ("Backend Functions","ACTIVE","delete-account v2, process-export v2, payment-methods v1, support-catalog v1"),
    ("Web Pages","Başarılı deploy","subscription-tracking.github.io/SubscriptTrack")
], [1.5,2.2,3.1])

doc.add_heading("İçindekiler", 1)
for x in ["1 Ürün ve platform kapsamı","2 Teknoloji yığını","3 Sistem mimarisi","4 Mobil kod yapısı","5 Supabase ve backend yapısı","6 Veri modeli ve güvenlik","7 Ana veri akışları","8 Bildirim ve kimlik akışları","9 Export destek ve servis kataloğu","10 Test ve dağıtım","11 Teknik riskler ve kabul listesi"]:
    bullet(doc, x)

doc.add_heading("1 Ürün ve Platform Kapsamı", 1)
doc.add_paragraph("Ürün amacı, kullanıcının Netflix, Spotify, SaaS, bulut depolama ve benzeri aboneliklerini tek merkezde izlemesi; yaklaşan yenilemelerden önce karar vermesi; iptal veya durdurma sonrasında tahmini tasarrufu görmesidir. Ürün finansal işlem başlatmaz ve kart/banka hesabı bağlamaz. Kullanıcı yalnızca ödeme yöntemi etiketi gibi metinsel bir isim saklayabilir.")
make_table(doc, ["Yüzey","Rol","Teknik not"], [
    ("iOS","Birincil mobil istemci","Flutter tek kod tabanı; native izin/deep-link katmanı"),
    ("Android","Birincil mobil istemci","Flutter; Android bildirim ve biyometri izinleri"),
    ("Web Pages","Önizleme ve auth dönüş yüzü","Flutter web; tam web dashboard değildir"),
    ("Supabase","Auth, PostgreSQL, RLS, Realtime, Functions","Proje ref: tdbljrojcmyjwchawfif")
], [1.2,2.4,3.2])

doc.add_heading("2 Teknoloji Yığını", 1)
make_table(doc, ["Katman","Teknoloji","Sorumluluk"], [
    ("Mobil","Flutter 3.44.8 / Dart 3.12.2","iOS + Android + web build"),
    ("Durum","Provider + ChangeNotifier","Controller durumları ve router refresh"),
    ("Navigasyon","go_router","Redirect guard ve deep-link rotaları"),
    ("Kimlik/veri","Supabase Flutter SDK","Auth, PostgREST, Realtime, Functions"),
    ("Yerel veri","flutter_secure_storage + cache","Hassas değerler ve çevrimdışı son veri"),
    ("Bildirim","flutter_local_notifications + timezone","Cihaz içi zamanlanmış bildirim"),
    ("Para","Money value object, minor unit","Floating point hatası olmadan hesap"),
    ("Backend","Supabase Edge Functions / Deno TypeScript","Delete, export, payment, support/catalog"),
    ("Veritabanı","PostgreSQL","UUID, NUMERIC, UTC TIMESTAMPTZ, RLS"),
    ("CI/CD","GitHub Actions + Pages","Web build/deploy ve kalite gate")
], [1.3,2.4,3.1])

doc.add_heading("3 Sistem Mimarisi", 1)
doc.add_paragraph("Mimari contract-first ve domain-first yaklaşımındadır. Mobil uygulama service role anahtarıyla veritabanına bağlanmaz; Auth oturum tokenı ile PostgREST/RPC veya kimlik doğrulamalı Edge Function sınırını kullanır. RLS, kullanıcının yalnızca kendi kayıtlarını görmesini ve değiştirmesini sağlar.")
code(doc, "Kullanıcı → Flutter mobil istemci → Supabase Auth / PostgREST / RPC / Realtime / Edge Functions → PostgreSQL")
doc.add_heading("3.1 Repository seçimi", 2)
code(doc, "API_BASE_URL tanımlıysa → ApiSubscriptionRepository\nSUPABASE_URL + SUPABASE_ANON_KEY varsa → SupabaseSubscriptionRepository\nAksi durumda → Local/Unavailable repository")
doc.add_paragraph("Bu seçim dart-define ile yapılır. Release APK ve Pages build’i Supabase URL ve anon key ile derlenmelidir. Anon key istemciye gömülebilir; service role key hiçbir mobil veya web build’ine konulmamalıdır.")
doc.add_heading("3.2 Realtime", 2)
doc.add_paragraph("SupabaseSubscriptionRepository.watchAll(userId), subscriptions tablosunu dinler. Controller gelen tam listeyi mevcut yerel listeyle birleştirir. 030_realtime_publication.sql ile subscriptions tablosu supabase_realtime publication’a eklenmiştir. İki cihazlı kabul testi gerçek oturum gerektirir.")

doc.add_heading("4 Mobil Kod Yapısı", 1)
code(doc, "mobile/lib/\n  app/                  router, theme, shell\n  core/                 config, domain, network, storage, services\n  features/auth/        Supabase/local auth ve ekranlar\n  features/subscriptions/  domain, repository, controller, ekranlar\n  features/dashboard/  toplamlar ve yaklaşan yenilemeler\n  features/calendar/   aylık takvim ve gün detayları\n  features/notifications/ in-app bildirim merkezi\n  features/stats/      ödeme geçmişi ve istatistikler\n  features/savings/    tasarruf olayları\n  features/settings/   profil, export, ödeme etiketleri, destek")
doc.add_paragraph("Ekranlar doğrudan veritabanı sorgusu yapmaz; controller üzerinden data source/repository kullanır. Controller’lar loading, error, offline ve sync durumlarını UI’a taşır.")
make_table(doc, ["Bileşen","Akış"], [
    ("AuthController","init, signIn, signUp, password reset, update email, delete account"),
    ("SubscriptionController","CRUD, status transition, offline queue, pagination, Realtime merge, ödeme yenileme"),
    ("SettingsController","Tema, para birimi, bildirim, timezone ve paymentMethods"),
    ("LocalStorage","Şifreli JSON cache; offline okuma ve kuyruk"),
    ("AppRouter","Splash → onboarding → auth → authenticated shell; recovery → reset-password")
], [1.7,5.1])

doc.add_heading("5 Supabase ve Backend Yapısı", 1)
make_table(doc, ["Fonksiyon / migration","Görev","Durum"], [
    ("delete-account v2","Auth kullanıcısını ve ilişkili veriyi silme","ACTIVE"),
    ("process-export v2","Canonical 11 alanlı CSV üretimi","ACTIVE"),
    ("payment-methods v1","Ödeme yöntemi etiketleri GET/PUT","ACTIVE"),
    ("support-catalog v1","Servis kataloğu GET, ticket POST","ACTIVE"),
    ("028_payment_methods_api","payment_methods API/RPC/RLS","Uygulandı"),
    ("029_support_and_service_catalog","support_tickets ve service_catalog","Uygulandı"),
    ("030_realtime_publication","subscriptions Realtime publication","Uygulandı")
], [2.0,3.8,1.0])
doc.add_paragraph("support_tickets endpoint’i JWT ister. Anon REST probe’da 401 görülmesi tablonun yokluğu değil, korumalı erişimin doğru çalıştığı anlamına gelir.")

doc.add_heading("6 Veri Modeli ve Güvenlik", 1)
doc.add_paragraph("Abonelik yaşam döngüsü fiziksel silme yerine active, paused, cancelled, expired ve archived durumlarıyla korunur. Fiziksel silme yalnızca yanlış kayıt veya hesap silme gibi özel akışlarda kullanılır. Para değerleri Money/minor unit yaklaşımıyla taşınır. Tarihler UTC saklanır, kullanıcı IANA timezone’ı ile gösterilir.")
make_table(doc, ["Katman","Uygulama"], [
    ("Auth","Supabase JWT; auth state stream; recovery mode"),
    ("RLS","user_id sahiplik filtresi"),
    ("Edge Functions","Service role yalnızca sunucu tarafında; JWT doğrulaması"),
    ("Secrets","URL ve anon key dart-define; service role client build’de yok"),
    ("İdempotency","Subscription RPC’leri ve kritik yazmalarda tekrar dayanıklılığı"),
    ("Yerel gizlilik","Secure storage; cache ve bildirim içeriklerinin cihaz sınırı")
], [1.8,5.0])

doc.add_heading("7 Ana Veri Akışları", 1)
doc.add_heading("7.1 Abonelik ekleme ve güncelleme", 2)
code(doc, "Form → SubscriptionController → repository → create/update_subscription_idempotent → PostgreSQL → cache + UI + notification reschedule")
doc.add_paragraph("Form servis kataloğundan seçim yapabilir veya manuel isim alabilir. Tutar, para birimi, billing cycle, yenileme tarihi, kategori, not, trial ve payment label normalize edilir.")
doc.add_heading("7.2 Ödendi ve yenilendi", 2)
doc.add_paragraph("Yenileme tarihi geçmiş veya bugün olan abonelikte kullanıcı Ödendi ve yenilendi aksiyonuna basar. Uygulama kart çekmez; manuel payment event kaydı oluşturur ve sonraki yenileme tarihini başlangıç gününü koruyarak ileri taşır.")
doc.add_heading("7.3 Offline", 2)
code(doc, "NetworkException → LocalStorage cache göster → isOffline banner\nYazma başarısız → OfflineMutationQueue → bağlantıda FIFO replay")
doc.add_heading("7.4 CSV", 2)
doc.add_paragraph("Yerel ve cloud export aynı canonical alanları kullanır: name, amount, currency, billing_cycle, next_renewal_date, category, status, notes, payment_method, trial_end_date, trial_price_after.")

doc.add_heading("8 Bildirim ve Kimlik Akışları", 1)
doc.add_heading("8.1 Bildirim", 2)
doc.add_paragraph("Bildirimler flutter_local_notifications ile cihaz üzerinde planlanır. Yenileme ve trial bitiş adayları oluşturulur, geçmiş tarihler atlanır, cihaz timezone’ı kullanılır ve en fazla 60 bekleyen bildirim tutulur. Ertele aksiyonu 30 dakika sonraya bildirim kurar. FCM/APNs tabanlı uzak push bulunmaz.")
doc.add_heading("8.2 Şifre sıfırlama", 2)
code(doc, "Forgot password → resetPasswordForEmail → e-mail link\n→ subscripttrack://auth-callback (native) veya Pages origin (web)\n→ passwordRecovery event → /reset-password → updateUser(password)")
doc.add_paragraph("E-posta doğrulaması zorunlu Supabase ayarında signUp oturumsuz dönerse uygulama kullanıcıyı yanlışlıkla authenticated işaretlemez; doğrulama e-postasını kontrol etmesi gerektiğini bildirir. E-posta değiştirme Auth updateUser(email) üzerinden yapılır.")
doc.add_paragraph("Gerçek mail teslimi, redirect URL ve native deep-link kabulü Supabase Dashboard ayarları ve gerçek cihaz/test hesabı gerektirir.")

doc.add_heading("9 Export Destek ve Servis Kataloğu", 1)
make_table(doc, ["Özellik","Mobil akış","Backend"], [
    ("Export","Export ekranı → Function → indirme","process-export v2"),
    ("Destek","Feedback → kategori + mesaj → ticket","support-catalog POST → support_tickets"),
    ("Servis kataloğu","SubscriptionForm remote presets; fallback","support-catalog GET → service_catalog"),
    ("Payment label","Ayarlar/form dropdown","payment-methods GET/PUT + replace RPC")
], [1.4,3.0,2.6])

doc.add_heading("10 Test ve Dağıtım", 1)
make_table(doc, ["Kontrol","Sonuç"], [
    ("flutter analyze --no-pub","Başarılı"),
    ("Controller/UI testleri","Başarılı; ilgili suite 36 test geçti"),
    ("Release APK","Başarılı; 59,9 MB"),
    ("GitHub Actions web deploy","Başarılı"),
    ("GitHub Pages","subscription-tracking.github.io/SubscriptTrack"),
    ("Supabase functions list","Dört fonksiyon ACTIVE"),
    ("Live inventory","Auth, storage ve tablolar erişilebilir; protected tablo 401 doğru")
], [3.2,3.8])
doc.add_paragraph("Smoke yardımcıları: backend/scripts/authenticated-subscription-smoke.mjs ve authenticated-feature-smoke.mjs. Gerçek test hesabı ve S31_SMOKE_CONFIRM=run olmadan uzak projede yazma yapmazlar.")

doc.add_heading("11 Teknik Riskler ve Kabul Listesi", 1)
make_table(doc, ["Risk / sınır","Önerilen aksiyon","Öncelik"], [
    ("SMTP ve Auth maili canlı kabul edilmedi","SMTP tanımla; gerçek inbox testi yap","Yüksek"),
    ("Native deep-link gerçek cihazda doğrulanmadı","Android ve iOS release cihaz testi","Yüksek"),
    ("İki cihaz Realtime kabulü yapılmadı","İki oturumla create/update/delete çalıştır","Yüksek"),
    ("Uzak push yok","Bu kararı koru veya FCM/APNs ayrı kapsam aç","Orta"),
    ("Pages URL Supabase redirect’e ekli olmalı","subscription-tracking.github.io/SubscriptTrack adresini Auth URL Configuration’a ekle","Yüksek"),
    ("APK binary GitHub’da tutuluyor","Uzun vadede Releases veya Git LFS’e taşı","Düşük")
], [2.5,3.8,.7])
for x in ["[ ] Supabase SMTP ve Auth URL ayarlarını doğrula","[ ] Gerçek mail ile şifre sıfırlama ve e-posta doğrulama","[ ] Android ve iOS deep-link reset testi","[ ] İki oturumla Realtime create/update/delete","[ ] Gerçek export indir ve CSV’yi tekrar içe aktar","[ ] Test hesabı ile hesap silme ve veri yokluğu kontrolü"]:
    bullet(doc, x)
doc.add_paragraph("Rapor kaynakları: SubscriptTrack-Documentation/PRODUCT.md, ARCHITECTURE.md, API.md, MOBILE_ARCHITECTURE.md; mobile/lib; backend/migrations; backend/supabase/functions; GitHub Actions ve canlı Supabase inventory.")

for p in doc.paragraphs:
    if p.style.name.startswith("Heading"): p.paragraph_format.keep_with_next = True
props = doc.core_properties; props.title = "SubscriptTrack Teknik ve Teknolojik Yapı Raporu"; props.subject = "Mobil uygulama, Supabase backend ve veri akışları"; props.author = "SubscriptTrack Teknik Ekip"
doc.save(OUT)
print(OUT)
