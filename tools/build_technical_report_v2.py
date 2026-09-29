from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

OUT = Path(r"D:\SubscriptTrack\outputs\reports\SubscriptTrack_Teknik_Yapi_ve_Ozellik_Envanteri.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY, BLUE, GRAY = "17365D", "2F75B5", "5B6573"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(t):
    tblPr = t._tbl.tblPr; bd = tblPr.first_child_found_in("w:tblBorders")
    if bd is None: bd = OxmlElement("w:tblBorders"); tblPr.append(bd)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = bd.find(qn("w:w:"+edge)) if False else bd.find(qn("w:"+edge))
        if e is None: e = OxmlElement("w:"+edge); bd.append(e)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "D9D9D9")

def cell(cell, text, bold=False, color="000000", size=8.5):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for i, htxt in enumerate(headers): cell(t.rows[0].cells[i], htxt, True, "FFFFFF", 9); shade(t.rows[0].cells[i], NAVY)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell(cells[i], val)
            if n % 2: shade(cells[i], "F6F9FC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(text)

def code(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("2F3B4A")

doc = Document()
sec = doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles = doc.styles; styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(10); styles["Normal"].font.color.rgb=RGBColor(35,42,52)
styles["Normal"].paragraph_format.space_after=Pt(6); styles["Normal"].paragraph_format.line_spacing=1.08
for name,size,color in [("Title",28,NAVY),("Heading 1",18,NAVY),("Heading 2",13,BLUE),("Heading 3",11,NAVY)]:
    s=styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6)
header=sec.header.paragraphs[0]; header.text="SUBSCRIPTRACK  |  TEKNOLOJİ VE ÖZELLİK ENVANTERİ"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY); r.bold=True
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run("SubscriptTrack • Güncel teknik görünüm"); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("SubscriptTrack Teknoloji ve Özellik Envanteri")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Ürünün hangi teknolojilerle çalıştığını ve hangi özelliklerin mevcut olduğunu açıklayan güncel rapor"); r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(GRAY)
doc.add_paragraph("")
table(doc,["Bilgi","Değer"],[("Güncelleme","22 Eylül 2026"),("Platform","iOS, Android ve Flutter Web önizlemesi"),("Mobil teknoloji","Flutter 3.44.8 / Dart 3.12.2"),("Backend","Supabase • PostgreSQL • Edge Functions"),("Güncel kaynak","main • commit 6e116b2"),("Release APK","SubscriptTrack-latest-release.apk")],[1.7,5.1])
doc.add_page_break()

doc.add_heading("Bu rapor neyi gösteriyor",1)
doc.add_paragraph("Bu doküman bir sprint planı veya test kabul listesi değildir. Ürünün teknik olarak nasıl kurulduğunu, kodun hangi bölümlere ayrıldığını, özelliklerin mevcut durumunu ve temel akışların nereden nereye gittiğini sade bir dille gösterir.")
doc.add_heading("Durum işaretleri",2)
table(doc,["Durum","Anlam"],[("Aktif","Kod, backend ve mobil bağlantısı mevcut; normal kullanım akışı çalışır."),("Kısmi","Kodun bir bölümü mevcut ancak gerçek hesap, e-posta, cihaz veya dış ayar doğrulaması eksik."),("Yok","Bu kapsamda uygulanmamış veya ürün kararı gereği kullanılmıyor.")],[1.3,5.5])

doc.add_heading("1 Teknoloji Yığını",1)
table(doc,["Katman","Kullanılan teknoloji","Ne için kullanılıyor"],[
("Mobil uygulama","Flutter 3.44.8, Dart 3.12.2","Tek kod tabanından Android ve iOS uygulaması"),
("UI ve durum","Material, Provider, ChangeNotifier","Ekranlar, form durumları, loading/error/offline state"),
("Navigasyon","go_router","Auth koruması, onboarding, deep-link ve ekran geçişleri"),
("Kimlik","Supabase Auth","Kayıt, giriş, e-posta doğrulama, şifre sıfırlama, e-posta değiştirme"),
("Veri API","Supabase PostgREST ve RPC","Abonelik CRUD, idempotent yazma ve kullanıcı verisi"),
("Veritabanı","Supabase PostgreSQL","Abonelikler, ödeme olayları, ayarlar, bildirimler, ticket’lar"),
("Yetkilendirme","JWT + Row Level Security","Kullanıcı verilerinin birbirinden ayrılması"),
("Canlı senkron","Supabase Realtime","Abonelik INSERT/UPDATE/DELETE değişikliklerini dinleme"),
("Sunucu fonksiyonları","Deno TypeScript Edge Functions","Hesap silme, export, ödeme etiketi, destek/katalog"),
("Yerel depolama","flutter_secure_storage + local cache","Hassas bilgiler ve çevrimdışı son veri"),
("Bildirim","flutter_local_notifications + timezone","Cihaz üzerinde yenileme/trial hatırlatmaları"),
("Para","Money value object ve minor unit","Floating point hatası olmadan tutar hesabı"),
("CI/CD","GitHub Actions ve GitHub Pages","Web build/deploy ve kalite kontrolleri")
],[1.4,2.5,3.1])

doc.add_heading("2 Kod Yapısı",1)
code(doc,"mobile/lib/\n  app/                  app shell, tema ve router\n  core/                 config, domain, network, storage, services\n  features/auth/        kayıt, giriş, şifre akışları\n  features/subscriptions/ abonelik domain, repository ve ekranlar\n  features/dashboard/  toplamlar ve yaklaşan yenilemeler\n  features/calendar/   aylık takvim\n  features/notifications/ bildirim merkezi\n  features/stats/      istatistik ve ödeme geçmişi\n  features/savings/    tasarruf olayları\n  features/settings/   profil, export, ödeme etiketleri, destek")
table(doc,["Kod bölümü","Ana dosyalar / sorumluluk"],[
("Auth","auth_controller.dart, supabase_auth_repository.dart, login/register/forgot/reset ekranları"),
("Abonelik","subscription_controller.dart, SupabaseSubscriptionRepository, SubscriptionForm"),
("Ayarlar","SettingsController, ProfileScreen, PaymentMethodsScreen, ExportDataScreen"),
("Bildirim","LocalNotificationService; planlama, izin, tıklama ve erteleme"),
("Yerel veri","LocalStorage, SecureStorage, OfflineMutationQueue"),
("Backend","backend/supabase/functions ve backend/migrations"),
("Test","mobile/test; controller, UI, parser, güvenlik ve veri akışı testleri")
],[1.5,5.5])

doc.add_heading("3 Sistem Nasıl Çalışıyor",1)
doc.add_paragraph("Mobil uygulama iki çalışma moduna göre repository seçer. SUPABASE_URL ve SUPABASE_ANON_KEY ile derlenmiş sürüm canlı Supabase’i kullanır. Bu değerler yoksa uygulama local/unavailable auth ve yerel veri davranışına düşer.")
code(doc,"Kullanıcı → Flutter ekranı → Controller → Repository/API → Supabase Auth/PostgREST/RPC/Function → PostgreSQL")
doc.add_heading("3.1 Abonelik ekleme",2)
code(doc,"SubscriptionForm → SubscriptionController.create() → create_subscription_idempotent RPC → subscriptions tablosu → cache + dashboard + bildirim planı")
doc.add_heading("3.2 Abonelik güncelleme",2)
code(doc,"EditSubscriptionScreen → update_subscription_idempotent RPC → yeni kayıt durumu → Realtime ve yerel liste güncellemesi")
doc.add_heading("3.3 Ödeme yapıldı akışı",2)
doc.add_paragraph("Kullanıcı yenileme günü geldiğinde detay ekranındaki Ödendi ve yenilendi düğmesine basar. Uygulama karttan para çekmez. Payment event oluşturur, aboneliğin sonraki yenilemesini hesaplar ve geçmişe ekler.")
doc.add_heading("3.4 Çevrimdışı davranış",2)
code(doc,"Ağ hatası → cache’ten son liste → isOffline görünümü\nBaşarısız yazma → OfflineMutationQueue → bağlantı gelince FIFO tekrar deneme")

doc.add_heading("4 Özellik Envanteri",1)
table(doc,["Özellik","Durum","Şu an ne var"],[
("Onboarding","Aktif","İlk açılış tanıtımı ve tamamlanma kaydı"),
("E-posta kayıt ve giriş","Aktif","Supabase Auth ile kayıt/giriş; local fallback mevcut"),
("Google / Apple giriş","Kısmi","Kod yolu mevcut; sağlayıcı ayarı ve canlı OAuth testi yok"),
("Şifremi unuttum","Kısmi","Form, Supabase reset e-postası ve reset ekranı var; mail/redirect canlı testi gerekiyor"),
("E-posta doğrulama","Kısmi","Oturumsuz sign-up sonucu doğru ele alınıyor; gerçek mail teslimi ayarı gerekiyor"),
("E-posta değiştirme","Aktif / canlı test bekliyor","Profile ekranı ve Auth updateUser akışı mevcut"),
("Abonelik ekleme","Aktif","Manuel form ve remote servis katalog önerileri"),
("Abonelik düzenleme","Aktif","Alanlar, status ve bildirim kuralları güncelleniyor"),
("Abonelik lifecycle","Aktif","Active, paused, cancelled, expired, archived durumları"),
("Abonelik silme","Aktif","Yanlış kayıt için fiziksel silme; normal yaşamda archive/cancel"),
("Pagination / load more","Aktif","API repository ve controller pagination altyapısı"),
("Dashboard","Aktif","Para birimi bazında aylık/yıllık toplam, yaklaşanlar ve durumlar"),
("Takvim","Aktif","Yenileme ve trial tarihleri gün görünümünde"),
("İstatistikler","Kısmi","Ödeme geçmişi ve grafikler mevcut; dönem seçici kapsamı sınırlı"),
("Manuel ödeme kaydı","Aktif","Ödeme ekleme, düzenleme ve silme"),
("Ödendi ve yenilendi","Aktif","Payment event + sonraki yenileme tarihini ilerletme"),
("Ödeme yöntemi etiketi","Aktif / canlı test bekliyor","Kart çekmeden label GET/PUT; Supabase Edge Function"),
("Yerel bildirimler","Aktif","Yenileme/trial planlama, izin, test, erteleme, deep-link"),
("Uzak push","Yok","FCM/APNs kullanılmıyor; ürün yerel bildirim kullanıyor"),
("Offline cache","Aktif","Son veri görünümü ve offline mutation queue"),
("Realtime çoklu cihaz","Kısmi","Kod ve publication mevcut; iki gerçek oturum testi bekliyor"),
("CSV export","Aktif / canlı test bekliyor","Yerel ve cloud canonical 11 alan uyumlu"),
("CSV import","Aktif","Aynı alan setini okuyup abonelik oluşturma"),
("Hesap silme","Aktif / canlı test bekliyor","delete-account v2 ve yerel temizlik"),
("Destek ticket","Aktif / canlı test bekliyor","support_tickets tablosu, ekran ve Function"),
("Servis kataloğu","Aktif","service_catalog tablosu, seed kayıtlar, mobil fallback"),
("Tema ve görünüm","Aktif","Light/dark, responsive ekran davranışları"),
("Biyometrik kilit","Kısmi","local_auth ve 5 dakika kilit süresi; cihaz testi gerekiyor"),
],[1.8,1.5,4.0])

doc.add_heading("5 Backend Özellikleri",1)
table(doc,["Canlı bileşen","Durum","İşlev"],[
("delete-account v2","ACTIVE","Auth kullanıcısını ve ilişkili verileri siler"),
("process-export v2","ACTIVE","Canonical CSV üretir"),
("payment-methods v1","ACTIVE","Kullanıcı ödeme etiketi listesini GET/PUT yapar"),
("support-catalog v1","ACTIVE","Servis kataloğu okur, destek ticket açar"),
("028 migration","Uygulandı","Payment methods API/RPC/RLS"),
("029 migration","Uygulandı","Support tickets ve service catalog"),
("030 migration","Uygulandı","subscriptions Realtime publication")
],[2.0,1.3,4.0])
doc.add_paragraph("Canlı Supabase inventory kontrolünde Auth, Storage ve ana tablolar erişilebilir. support_tickets anonim çağrıda 401 döner; bu, tablonun korumalı olduğu anlamına gelir.")

doc.add_heading("6 Bildirim Sistemi",1)
doc.add_paragraph("Bildirimler sunucudan gönderilen push değildir. Uygulama, abonelik kayıtlarını ve kullanıcı tercihlerini okuyarak bildirimi cihaz üzerinde planlar. Varsayılan saat 09:00’dur; 1, 3 veya 7 gün öncesi seçilebilir. Trial için 7, 3 ve 1 gün öncesi hatırlatma üretilir. Bildirime dokunmak ilgili abonelik detayını açar; Ertele 30 dakika sonrasına kurar. En fazla 60 bekleyen bildirim tutulur.")
doc.add_heading("7 Kimlik ve Şifre Akışı",1)
code(doc,"Giriş → ForgotPasswordScreen → resetPasswordForEmail() → e-posta linki\n→ web: Pages URL / native: subscripttrack://auth-callback\n→ passwordRecovery → ResetPasswordScreen → updateUser(password)")
doc.add_paragraph("Gerçek e-posta gönderimi Supabase Auth mail ayarına, web dönüşü ise GitHub Pages adresinin Supabase Auth URL Configuration’a eklenmesine bağlıdır.")

doc.add_heading("8 Dağıtım ve Güncel Sürüm",1)
table(doc,["Dağıtım yüzeyi","Güncel durum"],[
("Android release APK","outputs/apk/SubscriptTrack-latest-release.apk • 59,9 MB"),
("GitHub","main branch • commit 6e116b2"),
("GitHub Pages","https://subscription-tracking.github.io/SubscriptTrack/"),
("Pages workflow","Build ve deploy başarılı"),
("Supabase Functions","delete-account v2, process-export v2, payment-methods v1, support-catalog v1 ACTIVE")
],[2.2,5.0])

doc.add_heading("9 Kısa Teknik Sonuç",1)
doc.add_paragraph("SubscriptTrack’in mevcut ürün omurgası Flutter + Supabase üzerine kuruludur. Mobil uygulama ekranları controller/repository katmanından geçerek Supabase Auth, PostgREST, RPC, Realtime ve Edge Functions ile haberleşir. Abonelik takibi, ödeme geçmişi, ödeme etiketi, servis kataloğu, destek kaydı, CSV ve hesap yönetimi kod tarafında mevcuttur. Kısmi görünen maddeler ağırlıklı olarak gerçek e-posta, gerçek cihaz veya iki oturumla yapılması gereken canlı doğrulamalardır; ürünün temel kod bileşenleri bulunmaktadır.")

for p in doc.paragraphs:
    if p.style.name.startswith("Heading"): p.paragraph_format.keep_with_next = True
props=doc.core_properties; props.title="SubscriptTrack Teknoloji ve Özellik Envanteri"; props.subject="Teknoloji yığını, kod yapısı ve özellik durumu"; props.author="SubscriptTrack Teknik Ekip"
doc.save(OUT); print(OUT)

